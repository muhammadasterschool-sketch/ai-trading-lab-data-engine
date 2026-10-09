"""Platform package tests — workbook export (Phase E, Sheets A–G)."""

from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from data_engine.platform import (
    AuditLog,
    InstrumentRegistry,
    MarketWatchEntry,
    OpenPositionRecord,
    ProviderSymbolMapping,
    StrategyPerformanceRecord,
    TradeRecord,
    WorkbookExporter,
)
from data_engine.platform.news import NewsCategory, NewsIntelligenceService, Sentiment
from data_engine.schemas import AssetClass, Instrument

T0 = datetime(2026, 10, 10, 12, 0, tzinfo=UTC)


@pytest.fixture()
def registry():
    reg = InstrumentRegistry()
    gold = Instrument(
        symbol="XAU/USD", asset_class=AssetClass.METAL, base_asset="XAU", quote_asset="USD"
    )
    reg.register(
        gold,
        "csv:metals_daily",
        T0,
        provider_symbols=(
            ProviderSymbolMapping(provider="mt5:BrokerX", provider_symbol="XAUUSD"),
            ProviderSymbolMapping(provider="tradingview", provider_symbol="OANDA:XAUUSD"),
        ),
    )
    return reg


@pytest.fixture()
def trades():
    return [
        TradeRecord(
            trade_id="TRD-1",
            timestamp=T0,
            instrument_id="iid-gold",
            strategy_id="STR-1",
            strategy_version="1",
            broker="simulator",
            account_mode="paper",
            trade_mode="paper",
            direction="long",
            order_type="market",
            entry_price=2350.0,
            exit_price=2361.0,
            position_size=0.5,
            fees=1.2,
            realized_pl=4.3,
            status="closed",
            exit_reason="take_profit",
        ),
        TradeRecord(
            trade_id="TRD-2",
            timestamp=T0 + timedelta(hours=1),
            instrument_id="iid-gold",
            strategy_id="STR-1",
            strategy_version="1",
            broker="simulator",
            account_mode="paper",
            trade_mode="paper",
            direction="short",
            order_type="limit",
            entry_price=2361.0,
            position_size=0.25,
            status="open",
        ),
    ]


def test_xlsx_has_all_seven_sheets(tmp_path: Path, registry):
    exporter = WorkbookExporter(registry=registry)
    path = exporter.export_xlsx(tmp_path / "lab.xlsx", T0)
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "A Instrument Registry",
        "B Trade Log",
        "C Strategy Performance",
        "D Open Positions",
        "E Market Watch",
        "F System Audit Log",
        "G News Log",
    ]


def test_instrument_sheet_renders_registry_row(tmp_path: Path, registry):
    exporter = WorkbookExporter(registry=registry)
    path = exporter.export_xlsx(tmp_path / "lab.xlsx", T0)
    ws = load_workbook(path)["A Instrument Registry"]
    header = [c.value for c in ws[1]]
    assert "Internal Instrument ID" in header
    assert "Live Trading Enabled" in header
    row = [c.value for c in ws[3]]
    assert row[1] == "XAU/USD"
    assert "XAUUSD" in row[4]
    assert row[15] == "csv:metals_daily"
    # live trading column honestly shows Not Authorized
    live_idx = header.index("Live Trading Enabled")
    assert row[live_idx] == "Not Authorized"


def test_trade_log_rows_render_with_mode_separation(tmp_path: Path, registry, trades):
    exporter = WorkbookExporter(registry=registry, trades=trades)
    path = exporter.export_xlsx(tmp_path / "lab.xlsx", T0)
    ws = load_workbook(path)["B Trade Log"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    assert len(rows) == 2
    assert rows[0][0] == "TRD-1" and rows[1][0] == "TRD-2"
    # paper/demo/live separation is a first-class column
    header = [c.value for c in ws[1]]
    assert header[6] == "Account Mode" and header[7] == "Trade Mode"
    assert all(r[7] == "paper" for r in rows)


def test_duplicate_trade_ids_refused(tmp_path: Path, registry, trades):
    dup = trades[0].model_copy(deep=True)
    exporter = WorkbookExporter(registry=registry, trades=[trades[0], dup])
    with pytest.raises(ValueError, match="duplicate trade_id"):
        exporter.export_xlsx(tmp_path / "lab.xlsx", T0)


def test_audit_log_sheet_renders_events(tmp_path: Path, registry):
    log = AuditLog()
    log.record(T0, "registry", "instrument_registered", "gold registered", "iid-1")
    exporter = WorkbookExporter(registry=registry, audit_log=log)
    path = exporter.export_xlsx(tmp_path / "lab.xlsx", T0)
    ws = load_workbook(path)["F System Audit Log"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    assert len(rows) == 1
    assert rows[0][1] == "registry" and rows[0][2] == "instrument_registered"


def test_news_log_sheet_renders_items(tmp_path: Path, registry):
    svc = NewsIntelligenceService()
    svc.register(
        headline="Fed holds rates",
        source="reuters",
        url="https://example.com/fed",
        published_at=T0 - timedelta(hours=1),
        retrieved_at=T0,
        category=NewsCategory.MONETARY_POLICY,
        sentiment=Sentiment.NEUTRAL,
        affected_instruments=("iid-1",),
    )
    exporter = WorkbookExporter(registry=registry, news=svc.all_items())
    path = exporter.export_xlsx(tmp_path / "lab.xlsx", T0)
    ws = load_workbook(path)["G News Log"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    assert len(rows) == 1
    assert rows[0][1] == "Fed holds rates"
    assert rows[0][8] == "monetary_policy"


def test_open_positions_and_market_watch_render(tmp_path: Path, registry):
    pos = OpenPositionRecord(
        position_id="POS-1",
        mode="paper",
        instrument_id="iid-gold",
        strategy_id="STR-1",
        direction="long",
        entry_price=2350.0,
        current_price=2355.0,
        position_size=0.5,
        unrealized_pl=2.5,
        exposure=1177.5,
        last_update=T0,
    )
    watch = MarketWatchEntry(
        instrument_id="iid-gold",
        monitoring=True,
        last_price=2355.0,
        last_price_time=T0,
        timeframe="1d",
        data_fresh=True,
    )
    perf = StrategyPerformanceRecord(
        strategy_id="STR-1",
        strategy_version="1",
        instrument_id="iid-gold",
        trade_count=2,
        net_pl=4.3,
        win_rate=0.5,
        max_drawdown=-0.02,
        current_mode="paper",
        paper_status="running",
    )
    exporter = WorkbookExporter(
        registry=registry,
        open_positions=[pos],
        market_watch=[watch],
        performance=[perf],
    )
    path = exporter.export_xlsx(tmp_path / "lab.xlsx", T0)
    wb = load_workbook(path)
    assert list(wb["D Open Positions"].iter_rows(min_row=3, values_only=True))[0][1] == "paper"
    assert list(wb["E Market Watch"].iter_rows(min_row=3, values_only=True))[0][1] == "yes"
    assert list(wb["C Strategy Performance"].iter_rows(min_row=3, values_only=True))[0][3] == 2


def test_csv_export_is_byte_deterministic(tmp_path: Path, registry):
    e1 = WorkbookExporter(registry=registry)
    e2 = WorkbookExporter(registry=registry)
    d1 = tmp_path / "a"
    d2 = tmp_path / "b"
    files1 = e1.export_csv(d1, T0)
    files2 = e2.export_csv(d2, T0)
    assert len(files1) == 7
    for f1, f2 in zip(files1, files2):
        assert f1.read_bytes() == f2.read_bytes()


def test_csv_headers_match_mandate_columns(tmp_path: Path, registry):
    files = WorkbookExporter(registry=registry).export_csv(tmp_path, T0)
    by_name = {f.name: f for f in files}
    trade_header = by_name["trade_log.csv"].read_text().splitlines()[0]
    assert "Broker Order ID" in trade_header and "Audit Reference" in trade_header
    reg_header = by_name["instrument_registry.csv"].read_text().splitlines()[0]
    assert "Original MT5 Symbol" in reg_header and "TradingView Symbol" in reg_header


def test_last_sync_rendered_on_every_sheet(tmp_path: Path, registry):
    exporter = WorkbookExporter(registry=registry)
    path = exporter.export_xlsx(tmp_path / "lab.xlsx", T0)
    wb = load_workbook(path)
    for sheet in wb.sheetnames:
        assert wb[sheet]["A2"].value == "Last successful synchronization:"
        assert wb[sheet]["B2"].value == T0.isoformat()
